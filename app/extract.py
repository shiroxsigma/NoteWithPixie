"""Office ファイル（pptx/docx/xlsx/pdf）のテキスト抽出。関連ファイル参照の AI 文脈同梱用。

出力は Markdown 風テキストに統一する。AI に渡す文脈もエディタも Markdown 前提
なので、表や見出しをそのまま流せば追加の整形なしでプロンプトに載る。

各抽出ライブラリの import は関数内で行う。ライブラリが1つ欠けた環境でも
他の形式やアプリ本体が動き続けるようにするため（起動時 import で全滅させない）。
"""
from __future__ import annotations

from pathlib import Path

# フロントがハードコードせずに済むよう /api/settings 経由で公開される
SUPPORTED_EXTS: set[str] = {".pptx", ".docx", ".xlsx", ".pdf"}

# 抽出結果そのものの上限。llm.py 側にも 8000 字/ファイルのプロンプト上限があるが、
# あちらはコンテキスト長の都合。こちらは「数十MBの巨大ファイルを丸ごと文字列化して
# メモリとレスポンスを食い潰さない」ための独立した防波堤。
MAX_EXTRACT_CHARS = 50_000

# Office 系ファイルのサイズ上限。テキストの 1MB（files.MAX_BYTES）とは別枠。
# pptx は画像が埋め込まれるだけで数十MBになるのが普通で、抽出するのはテキスト
# だけなのでファイルサイズで弾く意味が薄い。それでも無制限は危険なので 50MB。
MAX_OFFICE_BYTES = 50_000_000

# xlsx の切り詰め閾値。シート全体を Markdown 表にするとデータ系シートで
# 容易に数万行になるため、AI 文脈として意味のある範囲に絞る。
XLSX_MAX_ROWS = 200
XLSX_MAX_COLS = 50

# 旧 Office 形式（.ppt/.doc/.xls）の実体は OLE2 コンテナ。拡張子だけ .pptx に
# 変えたファイルが持ち込まれることがあるので、先頭マジックで検出して案内する。
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _require(module: str, package: str, what: str):
    """抽出ライブラリを遅延 import する。無ければ日本語の ValueError に変換。

    ImportError をそのまま漏らすと API 層で 500 になる。ユーザーが直せる
    メッセージ（何を入れればよいか）にしてから投げ直す。"""
    import importlib
    try:
        return importlib.import_module(module)
    except ImportError:
        raise ValueError(
            f"{what}を読むには {package} が必要です（pipenv install で導入済みのはず）。"
        )


def _check_not_ole2(path: Path, old_ext: str) -> None:
    """旧形式（OLE2）の検出。zip ベースの新形式と間違えて渡された場合に、
    ライブラリの分かりにくい例外ではなく変換手順を案内するため先頭で調べる。"""
    with path.open("rb") as f:
        head = f.read(8)
    if head == _OLE2_MAGIC:
        raise ValueError(
            f"旧形式（{old_ext}）のファイルです。Office で開いて {path.suffix} 形式"
            "（新形式）への変換が必要です。"
        )


def _md_table(rows: list[list[str]]) -> str:
    """行列を Markdown の表に整形する。セル内の改行と | は表を壊すので潰す。"""
    if not rows:
        return ""
    def cell(v: str) -> str:
        return v.replace("|", "\\|").replace("\n", " ").strip()
    lines = ["| " + " | ".join(cell(c) for c in row) + " |" for row in rows]
    # 1行目をヘッダ扱いにする（Markdown の表はヘッダ必須のため）
    sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
    return "\n".join([lines[0], sep] + lines[1:])


# --- pptx ----------------------------------------------------------------------
def _pptx_shape_text(shape, parts: list[str]) -> None:
    """図形1つからテキストを拾う。グループ図形は中に図形を持つので再帰する。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            _pptx_shape_text(child, parts)
        return
    if shape.has_table:
        rows = [[c.text for c in row.cells] for row in shape.table.rows]
        parts.append(_md_table(rows))
        return
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)


def extract_pptx(path: Path) -> str:
    pptx = _require("pptx", "python-pptx", "PowerPoint ファイル")
    _check_not_ole2(path, ".ppt")
    try:
        prs = pptx.Presentation(str(path))
    except Exception as e:
        raise ValueError(f"PowerPoint ファイルを読み込めませんでした（破損の可能性）: {e}")
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts: list[str] = [f"## スライド {i}"]
        for shape in slide.shapes:
            _pptx_shape_text(shape, parts)
        # 発表者ノートは口頭で補う内容が書かれがちで、AI 文脈として価値が高い
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                parts.append(f"（ノート）{note}")
        out.append("\n\n".join(parts))
    return "\n\n".join(out)


# --- docx ----------------------------------------------------------------------
def extract_docx(path: Path) -> str:
    docx = _require("docx", "python-docx", "Word ファイル")
    _check_not_ole2(path, ".doc")
    import re

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        raise ValueError(f"Word ファイルを読み込めませんでした（破損の可能性）: {e}")
    out: list[str] = []
    # iter_inner_content は段落と表を文書内の出現順で返す（段落→表の順に
    # まとめて読むと文脈が入れ替わってしまうため、こちらを使う）
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            rows = [[c.text for c in row.cells] for row in block.rows]
            out.append(_md_table(rows))
            continue
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            # 見出しスタイルは # 見出しへ。英語名（Heading 1）と日本語 UI で
            # 付きうる名前（見出し 1）の両方を見る
            m = re.match(r"(?:Heading|見出し)\s*([1-4])", block.style.name or "")
            if m:
                out.append("#" * int(m.group(1)) + " " + text)
            else:
                out.append(text)
    return "\n\n".join(out)


# --- xlsx ----------------------------------------------------------------------
def extract_xlsx(path: Path) -> str:
    openpyxl = _require("openpyxl", "openpyxl", "Excel ファイル")
    _check_not_ole2(path, ".xls")
    try:
        # read_only: シート全体をメモリに展開しない。data_only: 数式でなく計算値を取る
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Excel ファイルを読み込めませんでした（破損の可能性）: {e}")
    try:
        out: list[str] = []
        for ws in wb.worksheets:
            parts: list[str] = [f"## {ws.title}"]
            rows: list[list[str]] = []
            truncated_rows = truncated_cols = False
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= XLSX_MAX_ROWS:
                    truncated_rows = True
                    break
                if len(row) > XLSX_MAX_COLS:
                    truncated_cols = True
                    row = row[:XLSX_MAX_COLS]
                rows.append(["" if v is None else str(v) for v in row])
            if rows:
                parts.append(_md_table(rows))
            if truncated_rows or truncated_cols:
                limits = []
                if truncated_rows:
                    limits.append(f"{XLSX_MAX_ROWS} 行")
                if truncated_cols:
                    limits.append(f"{XLSX_MAX_COLS} 列")
                parts.append(f"…（大きいため {'・'.join(limits)} で切り詰め）")
            out.append("\n\n".join(parts))
        return "\n\n".join(out)
    finally:
        wb.close()  # read_only モードはファイルハンドルを持ち続けるため明示的に閉じる


# --- pdf -----------------------------------------------------------------------
def extract_pdf(path: Path) -> str:
    pypdf = _require("pypdf", "pypdf", "PDF ファイル")
    try:
        reader = pypdf.PdfReader(str(path))
        out: list[str] = []
        for i, page in enumerate(reader.pages, 1):
            out.append(f"## ページ {i}\n\n{(page.extract_text() or '').strip()}")
        return "\n\n".join(out)
    except ValueError:
        raise
    except Exception as e:
        # pypdf は破損 PDF で独自例外を投げる。API 層で 400 にするため ValueError へ
        raise ValueError(f"PDF ファイルを読み込めませんでした（破損の可能性）: {e}")


# --- ディスパッチ ----------------------------------------------------------------
_HANDLERS = {
    ".pptx": extract_pptx,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".pdf": extract_pdf,
}


def extract_text(path: Path) -> str:
    """拡張子に応じてテキスト抽出する。先頭に `# ファイル名` を付けるのは、
    複数の参照を1つのプロンプトに並べたときにどのファイル由来か分かるようにするため。"""
    handler = _HANDLERS.get(path.suffix.lower())
    if handler is None:
        raise ValueError(f"テキスト抽出に対応していない形式です: {path.suffix}")
    text = f"# {path.name}\n\n{handler(path)}"
    if len(text) > MAX_EXTRACT_CHARS:
        text = text[:MAX_EXTRACT_CHARS] + f"\n…（{MAX_EXTRACT_CHARS} 文字で切り詰め）"
    return text
